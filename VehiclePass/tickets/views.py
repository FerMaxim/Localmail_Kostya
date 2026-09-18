from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.core.exceptions import PermissionDenied, ValidationError
from django.db.models import Q
from .models import Ticket, User
from .permissions import can_view_ticket, can_access_chat
from . import services

def landing_page(request):
    if request.user.is_authenticated:
        return redirect('tickets:dashboard')
    return render(request, 'tickets/landing.html')

@login_required
def dashboard(request):
    user = request.user
    base_qs = Ticket.objects.select_related('contractor').all().order_by('-created_at')

    # 1. Сценарий КОНТРАГЕНТА (Client Portal)
    if user.role == 'contractor':
        my_tickets = base_qs.filter(contractor=user)
        
        # KPI Счетчики для карточек
        total_count = my_tickets.count()
        in_work_count = my_tickets.filter(status__in=['new', 'bio_check']).count()
        action_required_count = my_tickets.filter(status='waiting_contractor').count()
        approved_count = my_tickets.filter(status='approved').count()
        rejected_count = my_tickets.filter(status='rejected').count()

        # Быстрый фильтр по табам (в 1 клик)
        current_tab = request.GET.get('tab', 'all')
        filtered_tickets = my_tickets
        if current_tab == 'action':
            filtered_tickets = my_tickets.filter(status='waiting_contractor')
        elif current_tab == 'work':
            filtered_tickets = my_tickets.filter(status__in=['new', 'bio_check'])
        elif current_tab == 'approved':
            filtered_tickets = my_tickets.filter(status='approved')
        elif current_tab == 'rejected':
            filtered_tickets = my_tickets.filter(status='rejected')

        # Лёгкий поиск по номеру ТС или грузу/месту
        vehicle_q = request.GET.get('vehicle', '').strip()
        if vehicle_q:
            filtered_tickets = filtered_tickets.filter(
                Q(vehicle_number__icontains=vehicle_q) |
                Q(cargo__icontains=vehicle_q) |
                Q(unloading_place__icontains=vehicle_q)
            )

        context = {
            'tickets': filtered_tickets,
            'total_count': total_count,
            'in_work_count': in_work_count,
            'action_required_count': action_required_count,
            'approved_count': approved_count,
            'rejected_count': rejected_count,
            'current_tab': current_tab,
            'vehicle': vehicle_q,
        }
        return render(request, 'tickets/dashboard_contractor.html', context)

    # 2. Сценарий СПЕЦИАЛИСТА ГЛОНASS (1-я линия модерации)
    elif user.role == 'glonass':
        # Счетчики очереди для диспетчера
        new_count = Ticket.objects.filter(status='new').count()
        bio_escalated_count = Ticket.objects.filter(status='bio_check').count()
        approved_count = Ticket.objects.filter(status='approved').count()
        
        tickets = base_qs
        
        # По умолчанию если нажали кнопку Входящие (или таб)
        tab = request.GET.get('tab', '')
        if tab == 'new':
            tickets = tickets.filter(status='new')
        elif tab == 'archive':
            tickets = tickets.exclude(status='new')

        # Профессиональные фильтры
        q = request.GET.get('q', '').strip()
        if q:
            tickets = tickets.filter(
                Q(comment__icontains=q) |
                Q(cargo__icontains=q) |
                Q(loading_place__icontains=q) |
                Q(unloading_place__icontains=q)
            )
            
        vehicle_q = request.GET.get('vehicle', '').strip()
        if vehicle_q:
            tickets = tickets.filter(vehicle_number__icontains=vehicle_q)

        contractor_q = request.GET.get('contractor', '').strip()
        if contractor_q:
            tickets = tickets.filter(
                Q(contractor__organization_name__icontains=contractor_q) |
                Q(org_name_override__icontains=contractor_q) |
                Q(contractor__username__icontains=contractor_q)
            )

        status_filters = request.GET.getlist('status')
        if status_filters and '' not in status_filters:
            tickets = tickets.filter(status__in=status_filters)

        contractor_names = list(set(
            User.objects.filter(role='contractor', organization_name__isnull=False)
            .exclude(organization_name='')
            .values_list('organization_name', flat=True)
        ))

        context = {
            'tickets': tickets,
            'new_count': new_count,
            'bio_escalated_count': bio_escalated_count,
            'approved_count': approved_count,
            'tab': tab,
            'q': q,
            'vehicle': vehicle_q,
            'contractor': contractor_q,
            'selected_statuses': status_filters,
            'contractor_names': contractor_names,
            'status_choices': Ticket.STATUS_CHOICES,
        }
        return render(request, 'tickets/dashboard_glonass.html', context)

    # 3. Сценарий ИНСПЕКТОРА БИОБЕЗОПАСНОСТИ (2-я линия контроля и карантина)
    elif user.role == 'bio_security':
        # Биобез видит только свои зоны ответственности
        bio_tickets = base_qs.filter(status__in=['bio_check', 'waiting_contractor', 'approved', 'rejected'])

        # Счетчики ветеринарно-санитарных задач
        review_count = Ticket.objects.filter(status='bio_check').count()
        waiting_count = Ticket.objects.filter(status='waiting_contractor').count()
        quarantine_count = Ticket.objects.filter(status='rejected', ban_until__isnull=False).count()

        tab = request.GET.get('tab', '')
        if tab == 'review':
            bio_tickets = bio_tickets.filter(status='bio_check')
        elif tab == 'waiting':
            bio_tickets = bio_tickets.filter(status='waiting_contractor')
        elif tab == 'quarantine':
            bio_tickets = bio_tickets.filter(status='rejected')

        q = request.GET.get('q', '').strip()
        if q:
            bio_tickets = bio_tickets.filter(
                Q(comment__icontains=q) | Q(loading_place__icontains=q) | Q(cargo__icontains=q)
            )

        vehicle_q = request.GET.get('vehicle', '').strip()
        if vehicle_q:
            bio_tickets = bio_tickets.filter(vehicle_number__icontains=vehicle_q)

        ban_until_filter = request.GET.get('ban_until')
        if ban_until_filter:
            bio_tickets = bio_tickets.filter(ban_until__gte=ban_until_filter, status='rejected')

        status_filters = request.GET.getlist('status')
        if status_filters and '' not in status_filters:
            bio_tickets = bio_tickets.filter(status__in=status_filters)

        context = {
            'tickets': bio_tickets,
            'review_count': review_count,
            'waiting_count': waiting_count,
            'quarantine_count': quarantine_count,
            'tab': tab,
            'q': q,
            'vehicle': vehicle_q,
            'ban_until': ban_until_filter,
            'selected_statuses': status_filters,
            'status_choices': [c for c in Ticket.STATUS_CHOICES if c[0] in ['bio_check', 'waiting_contractor', 'approved', 'rejected']],
        }
        return render(request, 'tickets/dashboard_bio.html', context)

    # 4. Fallback для Администратора/Суперпользователя без явной роли
    else:
        context = {
            'tickets': base_qs,
            'status_choices': Ticket.STATUS_CHOICES,
        }
        return render(request, 'tickets/dashboard.html', context)

@login_required
def ticket_detail(request, pk):
    ticket = get_object_or_404(
        Ticket.objects.select_related('contractor')
        .prefetch_related('attachments__uploaded_by', 'messages__sender', 'history__user'),
        pk=pk
    )
    
    if not can_view_ticket(request.user, ticket):
        messages.error(request, "У вас нет прав для просмотра данной заявки.")
        return redirect('tickets:dashboard')
        
    if request.method == 'POST':
        action = request.POST.get('action')
        try:
            if action == 'approve_glonass':
                services.approve_by_glonass(ticket, request.user)
                messages.success(request, "Заявка успешно согласована отделом ГЛОНАСС.")
            elif action == 'send_bio':
                services.escalate_to_bio(ticket, request.user)
                messages.warning(request, "Заявка переведена в отдел Биобезопасности для дополнительной проверки.")
            elif action == 'invalid_form':
                services.reject_by_glonass_invalid(ticket, request.user)
                messages.error(request, "Заявка отклонена по причине некорректного заполнения формы.")
            elif action == 'approve_bio':
                services.approve_by_bio(ticket, request.user)
                messages.success(request, "Заявка успешно согласована отделом Биобезопасности.")
            elif action == 'reject_bio':
                ban_date = request.POST.get('ban_until')
                services.ban_vehicle_by_bio(ticket, request.user, ban_date)
                messages.error(request, f"Въезд запрещен! ТС отправлено на карантин до {ban_date}.")
            elif action == 'revert_glonass':
                services.revert_to_new(ticket, request.user)
                messages.warning(request, "Заявка возвращена в статус Новая.")
            elif action == 'revert_bio':
                services.revert_to_bio_check(ticket, request.user)
                messages.warning(request, "Заявка возвращена на проверку ОАБ.")
            else:
                messages.error(request, "Неизвестное действие.")
        except (PermissionDenied, ValidationError) as e:
            messages.error(request, str(e))
        except Exception as e:
            messages.error(request, f"Ошибка при выполнении действия: {str(e)}")
            
        return redirect('tickets:ticket_detail', pk=ticket.id)

    can_view_chat, can_send_chat, chat_disable_reason = can_access_chat(request.user, ticket)
    
    context = {
        'ticket': ticket,
        'can_view_chat': can_view_chat,
        'can_send_chat': can_send_chat,
        'chat_disable_reason': chat_disable_reason,
    }
    return render(request, 'tickets/ticket_detail.html', context)

@login_required
def send_ticket_message(request, pk):
    ticket = get_object_or_404(Ticket, pk=pk)
    if request.method == 'POST':
        text = request.POST.get('text')
        file = request.FILES.get('attachment')
        try:
            services.send_chat_message(ticket, request.user, text=text, file=file)
            messages.success(request, "Сообщение отправлено в чат.")
        except (PermissionDenied, ValidationError) as e:
            messages.error(request, str(e))
        except Exception as e:
            messages.error(request, f"Ошибка при отправке: {str(e)}")
            
    return redirect('tickets:ticket_detail', pk=pk)

@login_required
def ticket_create(request):
    if request.user.role != 'contractor' and not request.user.is_superuser:
        messages.error(request, "Создавать заявки могут только Контрагенты.")
        return redirect('tickets:dashboard')
        
    if request.method == 'POST':
        try:
            loading_place = request.POST.get('loading_place')
            unloading_place = request.POST.get('unloading_place')
            cargo = request.POST.get('cargo')
            
            if not all([loading_place, unloading_place, cargo]):
                raise ValidationError("Место загрузки, место выгрузки и груз являются обязательными полями!")
                
            ticket = services.create_ticket(
                contractor=request.user,
                loading_place=loading_place,
                unloading_place=unloading_place,
                cargo=cargo,
                vehicle_number=request.POST.get('vehicle_number'),
                org_name_override=request.POST.get('org_name_override'),
                comment=request.POST.get('comment'),
                files=request.FILES.getlist('attachments')
            )
            messages.success(request, f"Заявка #{ticket.id} успешно создана и отправлена в отдел ГЛОНАСС.")
            return redirect('tickets:dashboard')
        except ValidationError as e:
            messages.error(request, str(e))
            
    return render(request, 'tickets/ticket_form.html')


import openpyxl
from django.http import HttpResponse
from datetime import datetime
from django.utils.timezone import make_aware

@login_required
def export_report_xlsx(request):
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    tickets = Ticket.objects.all().prefetch_related('history', 'contractor').order_by('created_at')

    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d')
            start_date = make_aware(start_date)
            tickets = tickets.filter(created_at__gte=start_date)
        except ValueError:
            pass

    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').replace(hour=23, minute=59, second=59)
            end_date = make_aware(end_date)
            tickets = tickets.filter(created_at__lte=end_date)
        except ValueError:
            pass

    wb = openpyxl.Workbook()
    
    # Sheet 1 (Summary)
    ws1 = wb.active
    ws1.title = "Сводка"
    
    period_str = f"С {start_date_str or 'начала'} по {end_date_str or 'конец'}"
    
    total_count = tickets.count()
    glonass_approved = 0
    glonass_rejected = 0
    oab_approved = 0
    oab_rejected = 0

    for t in tickets:
        oab_required = t.history.filter(
            Q(action_description__contains='Биобез') | Q(action_description__contains='ОАБ')
        ).exists()
        
        if oab_required:
            if t.status == 'approved':
                oab_approved += 1
            elif t.status == 'rejected':
                oab_rejected += 1
        else:
            if t.status == 'approved':
                glonass_approved += 1
            elif t.status in ['rejected', 'invalid_form']:
                glonass_rejected += 1

    ws1.append(["Период отчёта", period_str])
    ws1.append(["Всего ТС проверено", total_count])
    ws1.append(["Всего ТС согласовано ГЛОНАСС", glonass_approved])
    ws1.append(["Всего ТС не согласовано ГЛОНАСС", glonass_rejected])
    ws1.append(["Всего ТС согласовано ОАБ", oab_approved])
    ws1.append(["Всего ТС не согласовано ОАБ", oab_rejected])
    ws1.append(["Всего ТС принадлежат ТБ", "=COUNTIF('Доставка ТМЦ'!D:D, \"*ТБ*\")"])
    
    # Стилизуем немного колонки
    ws1.column_dimensions['A'].width = 45
    ws1.column_dimensions['B'].width = 25

    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    # Sheet 2 (Data)
    ws2 = wb.create_sheet(title="Доставка ТМЦ")
    headers = [
        "Дата проверки", "время проверки", "гос.номер ТС", "Контрагент",
        "откуда выехал", "место прибытия (КПК, название ПП, МПП)", "Груз (наименование, тип)",
        "результат проверки на пересечение известных зон АЧС", "Признак АЧС", "Необходимость согласования УВ/отдела по контролю биобезопасности",
        "решение вет службы (в случае пересечения известных зон АЧС)", "способ проверки маршрута ТС", "комментарии", "ФИО оператора"
    ]
    ws2.append(headers)
    
    for t in tickets:
        row_data = t.get_excel_data().split('\t')
        ws2.append(row_data)

    for col in range(1, 15):
        ws2.column_dimensions[get_column_letter(col)].width = 22

    if ws2.max_row > 1:
        tab = Table(displayName="TicketsTable", ref=f"A1:{get_column_letter(14)}{ws2.max_row}")
        style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        tab.tableStyleInfo = style
        ws2.add_table(tab)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=report_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'
    
    wb.save(response)
    return response

from django.contrib.auth import login
from .forms import ContractorRegistrationForm

def register_view(request):
    if request.user.is_authenticated:
        return redirect('tickets:dashboard')
        
    if request.method == 'POST':
        form = ContractorRegistrationForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.role = 'contractor'  # Все новые пользователи становятся Контрагентами
            user.save()
            login(request, user)
            messages.success(request, f"Регистрация прошла успешно! Добро пожаловать, {user.first_name} {user.last_name}.")
            return redirect('tickets:dashboard')
    else:
        form = ContractorRegistrationForm()
        
    return render(request, 'registration/register.html', {'form': form})

@user_passes_test(lambda u: u.is_superuser)
def personnel_management(request):
    if request.method == 'POST':
        user_id = request.POST.get('user_id')
        new_role = request.POST.get('role')
        if user_id and new_role in dict(User.ROLE_CHOICES):
            u = get_object_or_404(User, id=user_id)
            if not u.is_superuser:  # Предотвращение изменения роли самого себя/других суперадминов
                u.role = new_role
                u.save()
                messages.success(request, f"Роль пользователя {u.username} изменена на {u.get_role_display()}.")
            else:
                messages.error(request, "Невозможно изменить роль главного администратора.")
            return redirect('tickets:personnel')
            
    users = User.objects.exclude(is_superuser=True).order_by('-date_joined')
    return render(request, 'tickets/personnel.html', {'users': users, 'roles': User.ROLE_CHOICES})

